from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from io import BytesIO

import logging
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)


app = Flask(__name__)
CORS(app)


@app.route('/rapport', methods=['POST'])
def rapport(): 
    
    data = request.json
    logger.debug('Data received: %s', data)
    
    if data is None:
        return jsonify({"error": "No JSON data in the request"}), 400

    required_arguments = ['mandatory_fields_missing', 'mandatory_fields', 'file_fields']
    missing_arguments = [arg for arg in required_arguments if arg not in data]

    if missing_arguments:
        return jsonify({"error": f"Missing JSON argument(s): {', '.join(missing_arguments)}"}), 400
    
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "MandatoryField_Attendu"
    sheet["B1"] = "MandatoryField_Manquant"

    # Boucle pour remplir les cellules successives dans la colonne A avec les valeurs de data.mandatory_fields
    for i, field in enumerate(data['mandatory_fields'], start=2):
        sheet[f"A{i}"] = field
        if field in data['file_fields']:
            logger.debug('Field: %s', field)
            sheet[f"B{i}"] = field
        else:
            logger.debug('Field MANQUANT: %s', field)
            sheet[f"B{i}"] = "Manquant"
            
    last_row = 1 + len(data['mandatory_fields']) 
    last_row_to_add = 2 + len(data['mandatory_fields']) 

    sheet[f"A{last_row_to_add}"] =f'=COUNTA(A2:A{last_row})'
    sheet[f"B{last_row_to_add}"] =f'=COUNTIF(B2:B{last_row},"Manquant")'
    
    chart = BarChart()
    chart.type = "col"
    chart.title = "Field Missing vs Field Require"
    chart.y_axis.title = 'Number Field'
    y_values = Reference(sheet, min_col=1, min_row=last_row_to_add, max_col=1,max_row=last_row_to_add)
    x_values = Reference(sheet, min_col=1, min_row=1, max_col=2, max_row=1)
    chart.add_data(y_values, titles_from_data=True)
    chart.set_categories(x_values)
    sheet.add_chart(chart, "E3")
    
    # Save workbook to a bytes buffer
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    # Use Flask's send_file function to send the result
    return send_file(buffer, download_name='rapport_mandatory_fields.xlsx', as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')



@app.route('/rapport_mandatory_fields', methods=['POST'])
def rapport_mandatory_fields(): 
    
    data = request.json
    logger.debug('Data received: %s', data)
    
    if data is None:
        return jsonify({"error": "No JSON data in the request"}), 400

    required_arguments = ['mandatory_fields_missing', 'mandatory_fields', 'file_fields']
    missing_arguments = [arg for arg in required_arguments if arg not in data]

    if missing_arguments:
        return jsonify({"error": f"Missing JSON argument(s): {', '.join(missing_arguments)}"}), 400
    
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "MandatoryField_Attendu"
    sheet["B1"] = "MandatoryField_Manquant"

    # Boucle pour remplir les cellules successives dans la colonne A avec les valeurs de data.mandatory_fields
    for i, field in enumerate(data['mandatory_fields'], start=2):
        sheet[f"A{i}"] = field
        if field in data['file_fields']:
            logger.debug('Field: %s', field)
            sheet[f"B{i}"] = field
        else:
            logger.debug('Field MANQUANT: %s', field)
            sheet[f"B{i}"] = "Manquant"
            
    last_row = 1 + len(data['mandatory_fields']) 
    last_row_to_add = 2 + len(data['mandatory_fields']) 

    sheet[f"A{last_row_to_add}"] =f'=COUNTA(A2:A{last_row})'
    sheet[f"B{last_row_to_add}"] =f'=COUNTIF(B2:B{last_row},"Manquant")'
    
    chart = BarChart()
    chart.type = "col"
    chart.title = "Field Missing vs Field Require"
    chart.y_axis.title = 'Number Field'
    y_values = Reference(sheet, min_col=1, min_row=last_row_to_add, max_col=1,max_row=last_row_to_add)
    x_values = Reference(sheet, min_col=1, min_row=1, max_col=2, max_row=1)
    chart.add_data(y_values, titles_from_data=True)
    chart.set_categories(x_values)
    sheet.add_chart(chart, "E3")
    
    # Save workbook to a bytes buffer
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    # Use Flask's send_file function to send the result
    return send_file(buffer, download_name='rapport_mandatory_fields.xlsx', as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5007)
